#!/usr/bin/env python3
"""Merge the level and frequency lists into one candidate headword table.

Output: cache/wordlists.json
    { lemma: {"cefr": "B1"|None, "pos": [...], "lists": [...],
              "ngsl_rank": int|None, "wiki_rank": int|None, "wiki_count": int} }

The lists play different roles:

  CEFR-J       the level a Japanese learner is taught the word at (A1-B2 only).
               This is what orders the study course.
  NGSL/NAWL    corpus-based coverage lists. NAWL is the academic layer, which is
               where Kyoto-style expository passages live.
  TSL          TOEIC service list; kept as a tag, not as a course.
  enwiki       raw frequency over English Wikipedia. Expository register, so it
               ranks the "hard but real" words above B2 far better than a
               subtitle corpus would.
"""
import csv
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
RAW = os.path.join(HERE, "raw")
CACHE = os.path.join(HERE, "cache")

POS_MAP = {
    "noun": "n", "verb": "v", "adjective": "adj", "adverb": "adv",
    "preposition": "prep", "conjunction": "conj", "pronoun": "pron",
    "determiner": "det", "interjection": "intj", "number": "num",
    "auxiliary verb": "v", "modal verb": "v", "article": "det",
    "infinitive-marker": "prt", "particle": "prt", "exclamation": "intj",
}

WORD_RE = re.compile(r"^[a-z][a-z'\-]*$")


def norm_pos(p):
    if not p:
        return None
    return POS_MAP.get(p.strip().lower())


def read_cefrj(out):
    from openpyxl import load_workbook
    path = os.path.join(RAW, "CEFR-J Wordlist Ver1.6.xlsx")
    wb = load_workbook(path, read_only=True)
    ws = wb["ALL_sep"]
    n = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        head, pos, cefr = (row + (None, None, None))[:3]
        if not head or not cefr:
            continue
        for variant in str(head).split("/"):
            lemma = variant.strip().lower()
            if not WORD_RE.match(lemma):
                continue
            e = out.setdefault(lemma, blank())
            # Keep the lowest (easiest) level a headword is introduced at.
            if e["cefr"] is None or cefr < e["cefr"]:
                e["cefr"] = cefr
            p = norm_pos(pos)
            if p and p not in e["pos"]:
                e["pos"].append(p)
            if "cefrj" not in e["lists"]:
                e["lists"].append("cefrj")
            n += 1
    print(f"  CEFR-J: {n} rows")


def read_lemmatized(path, tag, out):
    """NGSL/NAWL/TSL 'for teaching' files: headword,inflection,inflection,..."""
    n = 0
    with open(path, encoding="utf-8-sig", errors="replace") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            head = line.split(",")[0].strip().lower()
            if not WORD_RE.match(head):
                continue
            e = out.setdefault(head, blank())
            if tag not in e["lists"]:
                e["lists"].append(tag)
            n += 1
    print(f"  {tag}: {n} headwords")


def read_ngsl_stats(out):
    path = os.path.join(RAW, "NGSL_12_stats.csv")
    with open(path, encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            lemma = (row.get("Lemma") or "").strip().lower()
            if not WORD_RE.match(lemma):
                continue
            e = out.setdefault(lemma, blank())
            try:
                e["ngsl_rank"] = int(row["SFI Rank"])
            except (KeyError, ValueError, TypeError):
                pass


def read_nawl_definitions(out):
    """NAWL ships an English definition and a Japanese translation per word."""
    path = os.path.join(RAW, "NAWL_12_with_en_definitions.csv")
    n = 0
    with open(path, encoding="utf-8-sig", errors="replace") as f:
        for row in csv.DictReader(f):
            lemma = (row.get("Meanings") or "").strip().lower()
            if not WORD_RE.match(lemma):
                continue
            e = out.setdefault(lemma, blank())
            e["nawl_def"] = (row.get("English Definition") or "").strip()
            e["nawl_ja"] = (row.get("J Translation") or "").strip()
            n += 1
    print(f"  NAWL definitions: {n}")


def read_wiki_frequency(out, keep_top=60000):
    path = os.path.join(RAW, "enwiki-word-frequency.txt")
    rank = 0
    with open(path, encoding="utf-8", errors="replace") as f:
        for line in f:
            parts = line.split()
            if len(parts) != 2:
                continue
            word, count = parts[0].lower(), int(parts[1])
            if not WORD_RE.match(word) or len(word) < 2:
                continue
            rank += 1
            if rank > keep_top:
                break
            e = out.setdefault(word, blank())
            e["wiki_rank"] = rank
            e["wiki_count"] = count
    print(f"  enwiki: kept top {min(rank, keep_top)}")


def blank():
    return {"cefr": None, "pos": [], "lists": [], "ngsl_rank": None,
            "wiki_rank": None, "wiki_count": 0, "nawl_def": "", "nawl_ja": ""}


def main():
    os.makedirs(CACHE, exist_ok=True)
    out = {}
    print("reading word lists")
    read_cefrj(out)
    read_lemmatized(os.path.join(RAW, "NGSL_12_lemmatized_for_teaching.csv"), "ngsl", out)
    read_lemmatized(os.path.join(RAW, "NAWL_12_lemmatized_for_teaching.csv"), "nawl", out)
    read_lemmatized(os.path.join(RAW, "TSL_12_lemmatized_for_teaching.csv"), "tsl", out)
    read_ngsl_stats(out)
    read_nawl_definitions(out)
    read_wiki_frequency(out)

    graded = sum(1 for e in out.values() if e["lists"] and e["lists"] != ["tsl"])
    print(f"total candidates: {len(out)} ({graded} on a graded list)")
    with open(os.path.join(CACHE, "wordlists.json"), "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False)
    return 0


if __name__ == "__main__":
    sys.exit(main())
